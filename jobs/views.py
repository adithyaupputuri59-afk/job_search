from rest_framework.permissions import IsAuthenticated
from .models import Application
from .application_serializers import ApplicationSerializer
import io
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
from django.core.paginator import Paginator
from django.http import HttpResponse
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes, parser_classes
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response

from .models import Job
from .serializers import JobSerializer
from .permissions import IsRecruiterOrReadOnly

@api_view(["GET", "POST", "PUT", "PATCH", "DELETE"])
@permission_classes([IsRecruiterOrReadOnly])
def job_list(request, pk=None):

    if request.method == "GET":
        if pk is not None:
            try:
                job = Job.objects.get(pk=pk)
            except Job.DoesNotExist:
                return Response(
                    {"error": "Job not found"},
                    status=status.HTTP_404_NOT_FOUND
                )

            serializer = JobSerializer(job)
            return Response(serializer.data)

        jobs = Job.objects.all()

        title = request.query_params.get("title")
        location = request.query_params.get("location")

        if title:
            jobs = jobs.filter(title__icontains=title)

        if location:
            jobs = jobs.filter(location__icontains=location)

        page_number = request.query_params.get("page", 1)
        page_size = request.query_params.get("page_size", 10)

        try:
            page_size = int(page_size)
        except (TypeError, ValueError):
            page_size = 10

        page_size = max(1, min(page_size, 100))

        paginator = Paginator(jobs, page_size)
        page_obj = paginator.get_page(page_number)

        serializer = JobSerializer(
            page_obj.object_list,
            many=True
        )

        return Response({
            "total_jobs": paginator.count,
            "total_pages": paginator.num_pages,
            "current_page": page_obj.number,
            "next_page": (
                page_obj.next_page_number()
                if page_obj.has_next() else None
            ),
            "previous_page": (
                page_obj.previous_page_number()
                if page_obj.has_previous() else None
            ),
            "results": serializer.data
        })

    elif request.method == "POST":
        if isinstance(request.data, list):
            serializer = JobSerializer(
                data=request.data,
                many=True
            )
        else:
            serializer = JobSerializer(data=request.data)

        if serializer.is_valid():
            serializer.save()
            return Response(
                serializer.data,
                status=status.HTTP_201_CREATED
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )

    elif request.method in ["PUT", "PATCH"]:
        if pk is None:
            return Response(
                {"error": "Job ID is required"},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            job = Job.objects.get(pk=pk)
        except Job.DoesNotExist:
            return Response(
                {"error": "Job not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        serializer = JobSerializer(
            job,
            data=request.data,
            partial=(request.method == "PATCH")
        )

        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )

    elif request.method == "DELETE":
        if pk is None:
            return Response(
                {"error": "Job ID is required"},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            job = Job.objects.get(pk=pk)
        except Job.DoesNotExist:
            return Response(
                {"error": "Job not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        job.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

    return Response(
        {"error": "Method not supported"},
        status=status.HTTP_405_METHOD_NOT_ALLOWED
    )


@api_view(["DELETE"])
@permission_classes([IsRecruiterOrReadOnly])
def delete_all_jobs(request):
    Job.objects.all().delete()

    return Response(
        {"message": "All jobs deleted successfully"},
        status=status.HTTP_200_OK
    )


@api_view(["POST"])
@permission_classes([IsRecruiterOrReadOnly])
@parser_classes([MultiPartParser, FormParser])
def upload_excel(request):
    """
    Bulk upload jobs using an Excel file (.xlsx or .xls).
    Requires recruiter credentials (staff user).
    """
    if "file" not in request.FILES:
        return Response(
            {
                "status": "error",
                "message": "No file uploaded. Please upload an Excel file (.xlsx or .xls) using the 'file' key in form-data."
            },
            status=status.HTTP_400_BAD_REQUEST
        )

    excel_file = request.FILES["file"]

    if not (excel_file.name.endswith(".xlsx") or excel_file.name.endswith(".xls")):
        return Response(
            {
                "status": "error",
                "message": "Invalid file format. Please upload a valid Excel file (.xlsx or .xls)."
            },
            status=status.HTTP_400_BAD_REQUEST
        )

    try:
        df = pd.read_excel(excel_file)
    except Exception as e:
        return Response(
            {
                "status": "error",
                "message": f"Unable to read Excel file: {str(e)}"
            },
            status=status.HTTP_400_BAD_REQUEST
        )

    if df.empty:
        return Response(
            {
                "status": "error",
                "message": "The uploaded Excel file contains no data rows."
            },
            status=status.HTTP_400_BAD_REQUEST
        )

    # Normalize column names (lowercase, stripped, spaces to underscores)
    df.columns = [str(c).strip().lower().replace(" ", "_") for c in df.columns]

    required_fields = ["title", "company", "location", "salary", "description"]
    missing_fields = [f for f in required_fields if f not in df.columns]
    if missing_fields:
        return Response(
            {
                "status": "error",
                "message": f"Missing required columns in Excel: {', '.join(missing_fields)}. Expected: title, company, location, salary, description, job_type, experience"
            },
            status=status.HTTP_400_BAD_REQUEST
        )

    df = df.where(pd.notnull(df), None)
    records = df.to_dict(orient="records")

    for item in records:
        if "job_type" not in item or item["job_type"] is None:
            item["job_type"] = "Full-time"
        if "experience" not in item or item["experience"] is None:
            item["experience"] = "Fresher"

    serializer = JobSerializer(data=records, many=True)
    if serializer.is_valid():
        saved_jobs = serializer.save()
        return Response(
            {
                "status": "success",
                "message": f"Successfully uploaded {len(saved_jobs)} jobs from Excel file.",
                "total_uploaded": len(saved_jobs),
                "jobs": serializer.data
            },
            status=status.HTTP_201_CREATED
        )

    return Response(
        {
            "status": "error",
            "message": "Validation failed on some rows in the Excel file.",
            "errors": serializer.errors
        },
        status=status.HTTP_400_BAD_REQUEST
    )


@api_view(["GET"])
@permission_classes([IsRecruiterOrReadOnly])
def download_sample_excel(request):
    """
    Download a sample Excel template (.xlsx) for bulk job uploads.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jobs_Template"

    headers = ["title", "company", "location", "salary", "description", "job_type", "experience"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sample_rows = [
        ["Python Developer", "TCS", "Hyderabad", 40000, "Develop Python applications and maintain backend services.", "Full-time", "Fresher"],
        ["Django Developer", "Wipro", "Chennai", 45000, "Build REST APIs and backend applications using Django.", "Full-time", "Fresher"],
        ["Data Analyst", "Accenture", "Hyderabad", 38000, "Analyze datasets and prepare reports using SQL and Python.", "Full-time", "0-1 years"],
        ["Frontend Developer", "Tech Mahindra", "Pune", 32000, "Create responsive web pages using HTML, CSS, and JavaScript.", "Full-time", "Fresher"],
        ["DevOps Engineer", "IBM", "Bangalore", 50000, "Work with Linux, deployment pipelines, and cloud infrastructure.", "Full-time", "1-2 years"]
    ]

    for row in sample_rows:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="sample_jobs_template.xlsx"'
    return response
@api_view(["GET", "POST"])
@permission_classes([IsAuthenticated])
def application_list(request):

    if request.method == "GET":
        applications = Application.objects.filter(
            applicant=request.user
        )

        serializer = ApplicationSerializer(
            applications,
            many=True
        )

        return Response(serializer.data)

    job_id = request.data.get("job")

    if not job_id:
        return Response(
            {"error": "Job ID is required"},
            status=status.HTTP_400_BAD_REQUEST
        )

    try:
        job = Job.objects.get(pk=job_id)
    except (Job.DoesNotExist, ValueError, TypeError):
        return Response(
            {"error": "Job not found"},
            status=status.HTTP_404_NOT_FOUND
        )

    if Application.objects.filter(
        applicant=request.user,
        job=job
    ).exists():
        return Response(
            {"error": "You have already applied for this job"},
            status=status.HTTP_400_BAD_REQUEST
        )

    application = Application.objects.create(
        applicant=request.user,
        job=job
    )

    serializer = ApplicationSerializer(application)

    return Response(
        serializer.data,
        status=status.HTTP_201_CREATED
    )